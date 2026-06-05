"""DOCX input parser using python-docx."""

from __future__ import annotations

import os

from server.parsers.base import BaseParser


class DocxParser(BaseParser):
    """Extract text from Word documents."""

    def parse(self, source: str) -> str:
        """Extract text from a DOCX file.

        Args:
            source: Path to the DOCX file.

        Returns:
            Extracted text content with paragraph structure preserved.
        """
        from docx import Document

        if not os.path.exists(source):
            raise FileNotFoundError(f"DOCX file not found: {source}")

        doc = Document(source)
        paragraphs: list[str] = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                paragraphs.append("")
                continue

            # Preserve heading structure
            if para.style and para.style.name.startswith("Heading"):
                try:
                    level = int(para.style.name.replace("Heading", "").strip())
                except ValueError:
                    level = 1
                prefix = "#" * level
                paragraphs.append(f"{prefix} {text}")
            else:
                paragraphs.append(text)

        # Also extract text from tables
        for table in doc.tables:
            paragraphs.append("")
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                paragraphs.append(" | ".join(cells))

        return "\n".join(paragraphs)

    def can_handle(self, source: str, source_type: str) -> bool:
        if source_type == "docx":
            return True
        if source_type == "auto" and os.path.isfile(source):
            return source.lower().endswith(".docx")
        return False

"""PPTX rendering engine — orchestrates theme + layout to produce final .pptx."""

from __future__ import annotations

import logging
import os
from datetime import datetime

from pptx import Presentation
from pptx.util import Inches

from server.models.schema import DeckPlan, SlideType
from server.renderer.themes import get_theme
from server.renderer.themes.base import Theme

logger = logging.getLogger(__name__)

# Map slide types to their render functions
_LAYOUT_MAP = {}


def _get_layout_map():
    """Lazy-load layout map to avoid circular imports."""
    global _LAYOUT_MAP
    if not _LAYOUT_MAP:
        from server.renderer.layouts import (
            render_agenda,
            render_chart,
            render_closing,
            render_content,
            render_key_takeaway,
            render_quote,
            render_section,
            render_table,
            render_title,
            render_two_column,
        )
        _LAYOUT_MAP = {
            SlideType.TITLE: render_title,
            SlideType.AGENDA: render_agenda,
            SlideType.SECTION: render_section,
            SlideType.CONTENT: render_content,
            SlideType.TWO_COLUMN: render_two_column,
            SlideType.CHART: render_chart,
            SlideType.TABLE: render_table,
            SlideType.QUOTE: render_quote,
            SlideType.KEY_TAKEAWAY: render_key_takeaway,
            SlideType.CLOSING: render_closing,
        }
    return _LAYOUT_MAP


def render_deck(deck: DeckPlan, output_path: str = "") -> str:
    """Render a complete deck plan to a PPTX file.

    Args:
        deck: Validated deck plan.
        output_path: Output file path. If empty, auto-generates one.

    Returns:
        Absolute path to the generated PPTX file.
    """
    theme = get_theme(deck.theme)
    layout_map = _get_layout_map()

    # Create presentation with widescreen dimensions
    prs = Presentation()
    prs.slide_width = Inches(theme.layout.slide_width)
    prs.slide_height = Inches(theme.layout.slide_height)

    # Use blank layout (index 6 is typically blank)
    blank_layout = prs.slide_layouts[6]

    total_slides = len(deck.slides)

    for i, slide_content in enumerate(deck.slides):
        slide = prs.slides.add_slide(blank_layout)
        page_num = i + 1

        render_fn = layout_map.get(slide_content.slide_type)
        if render_fn is None:
            logger.warning(f"Unknown slide type: {slide_content.slide_type}, skipping slide {page_num}")
            continue

        try:
            render_fn(slide, slide_content, theme, page_num, total_slides)
        except Exception as e:
            logger.error(f"Error rendering slide {page_num} ({slide_content.slide_type}): {e}")
            raise

        # Add speaker notes
        if slide_content.speaker_notes:
            notes_slide = slide.notes_slide
            notes_tf = notes_slide.notes_text_frame
            notes_tf.text = slide_content.speaker_notes

    # Determine output path
    if not output_path:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_dir = os.path.join(os.path.expanduser("~"), "ppt-cc-output")
        os.makedirs(output_dir, exist_ok=True)
        output_path = os.path.join(output_dir, f"presentation_{timestamp}.pptx")

    # Ensure output directory exists
    output_dir = os.path.dirname(output_path)
    if output_dir:
        os.makedirs(output_dir, exist_ok=True)

    prs.save(output_path)
    logger.info(f"Saved presentation to: {output_path}")

    return os.path.abspath(output_path)

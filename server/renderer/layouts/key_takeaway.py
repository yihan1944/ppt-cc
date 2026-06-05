"""Key takeaway slide."""

from __future__ import annotations

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from server.models.schema import SlideContent
from server.renderer.components import add_title_bar, add_page_number
from server.renderer.themes.base import Theme


def render_key_takeaway(slide, content: SlideContent, theme: Theme, page_num: int, total: int) -> None:
    """Render a key takeaway slide."""
    add_title_bar(slide, theme)
    add_page_number(slide, theme, page_num, total)

    layout = theme.layout
    r_primary, g_primary, b_primary = theme.hex_to_rgb(theme.colors.primary)
    r_muted, g_muted, b_muted = theme.hex_to_rgb(theme.colors.text_muted)

    # "KEY TAKEAWAY" label
    label_box = slide.shapes.add_textbox(
        Inches(layout.margin_left + 1.0),
        Inches(2.4),
        Inches(layout.content_width - 2.0),
        Inches(0.5),
    )
    lf = label_box.text_frame
    lp = lf.paragraphs[0]
    lp.alignment = PP_ALIGN.CENTER
    lr = lp.add_run()
    lr.text = "KEY TAKEAWAY"
    lr.font.size = Pt(11)
    lr.font.bold = True
    lr.font.name = "Calibri"
    lr.font.color.rgb = RGBColor(r_muted, g_muted, b_muted)

    # Key point text
    key_box = slide.shapes.add_textbox(
        Inches(layout.margin_left + 1.0),
        Inches(3.0),
        Inches(layout.content_width - 2.0),
        Inches(2.5),
    )
    ktf = key_box.text_frame
    ktf.word_wrap = True
    kp = ktf.paragraphs[0]
    kp.alignment = PP_ALIGN.CENTER
    kr = kp.add_run()
    kr.text = content.key_point
    kr.font.size = Pt(22)
    kr.font.bold = True
    kr.font.name = "Calibri Light"
    kr.font.color.rgb = RGBColor(r_primary, g_primary, b_primary)

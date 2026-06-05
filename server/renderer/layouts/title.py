"""Title slide — clean, left-aligned, generous whitespace."""

from __future__ import annotations

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from server.models.schema import SlideContent
from server.renderer.components import set_slide_background
from server.renderer.themes.base import Theme


def render_title(slide, content: SlideContent, theme: Theme, page_num: int, total: int) -> None:
    """Render a clean title slide.

    Left-aligned title, generous whitespace, no decoration.
    """
    set_slide_background(slide, theme, theme.colors.primary)

    layout = theme.layout
    r_light, g_light, b_light = theme.hex_to_rgb(theme.colors.text_light)
    r_accent, g_accent, b_accent = theme.hex_to_rgb(theme.colors.accent)

    # Title — large, left-aligned, white on dark
    txBox = slide.shapes.add_textbox(
        Inches(layout.margin_left),
        Inches(2.2),
        Inches(layout.content_width),
        Inches(1.2),
    )
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = content.title
    run.font.size = Pt(44)
    run.font.bold = True
    run.font.name = theme.fonts.title
    run.font.color.rgb = RGBColor(r_light, g_light, b_light)

    # Subtitle — muted, left-aligned, below title
    if content.subtitle:
        parts = content.subtitle.split("·")
        y_sub = 3.6

        if len(parts) >= 2:
            # Role line
            txBox2 = slide.shapes.add_textbox(
                Inches(layout.margin_left),
                Inches(y_sub),
                Inches(layout.content_width),
                Inches(0.4),
            )
            tf2 = txBox2.text_frame
            p2 = tf2.paragraphs[0]
            p2.alignment = PP_ALIGN.LEFT
            r2 = p2.add_run()
            r2.text = parts[0].strip()
            r2.font.size = Pt(16)
            r2.font.name = "Calibri"
            r2.font.color.rgb = RGBColor(r_accent, g_accent, b_accent)

            # Specialty line
            txBox3 = slide.shapes.add_textbox(
                Inches(layout.margin_left),
                Inches(y_sub + 0.45),
                Inches(layout.content_width),
                Inches(0.4),
            )
            tf3 = txBox3.text_frame
            p3 = tf3.paragraphs[0]
            p3.alignment = PP_ALIGN.LEFT
            r3 = p3.add_run()
            r3.text = " · ".join(p.strip() for p in parts[1:])
            r3.font.size = Pt(14)
            r3.font.name = "Calibri"
            r3.font.color.rgb = RGBColor(r_accent, g_accent, b_accent)
        else:
            txBox2 = slide.shapes.add_textbox(
                Inches(layout.margin_left),
                Inches(y_sub),
                Inches(layout.content_width),
                Inches(0.4),
            )
            tf2 = txBox2.text_frame
            p2 = tf2.paragraphs[0]
            p2.alignment = PP_ALIGN.LEFT
            r2 = p2.add_run()
            r2.text = content.subtitle
            r2.font.size = Pt(16)
            r2.font.name = "Calibri"
            r2.font.color.rgb = RGBColor(r_accent, g_accent, b_accent)

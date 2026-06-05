"""Agenda slide — structured with title bar, centered list."""

from __future__ import annotations

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from server.models.schema import SlideContent
from server.renderer.components import add_title_bar, add_title_text, add_page_number
from server.renderer.themes.base import Theme


def render_agenda(slide, content: SlideContent, theme: Theme, page_num: int, total: int) -> None:
    """Render an agenda slide."""
    add_title_bar(slide, theme)
    title_bottom = add_title_text(slide, theme, content.title or "Agenda")
    add_page_number(slide, theme, page_num, total)

    layout = theme.layout
    r_num, g_num, b_num = theme.hex_to_rgb(theme.colors.secondary)
    r_text, g_text, b_text = theme.hex_to_rgb(theme.colors.text_dark)

    items = content.agenda_items
    num_items = len(items)

    bottom_limit = layout.slide_height - 1.0
    available_height = bottom_limit - title_bottom
    item_height = max(0.95, available_height / max(num_items, 1))
    if num_items * item_height > available_height:
        item_height = available_height / num_items
    y_start = title_bottom + (available_height - num_items * item_height) / 2

    for i, item_text in enumerate(items):
        y = y_start + i * item_height
        text_y = y + (item_height - 0.5) / 2 + 0.06

        # Number
        num_box = slide.shapes.add_textbox(
            Inches(layout.margin_left),
            Inches(text_y),
            Inches(0.5),
            Inches(0.5),
        )
        nf = num_box.text_frame
        nf.word_wrap = False
        np = nf.paragraphs[0]
        np.alignment = PP_ALIGN.RIGHT
        nr = np.add_run()
        nr.text = str(i + 1)
        nr.font.size = Pt(18)
        nr.font.bold = True
        nr.font.name = "Calibri Light"
        nr.font.color.rgb = RGBColor(r_num, g_num, b_num)

        # Item text
        txt_box = slide.shapes.add_textbox(
            Inches(layout.margin_left + 0.8),
            Inches(text_y),
            Inches(layout.content_width - 0.8),
            Inches(0.5),
        )
        tf = txt_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = item_text
        run.font.size = Pt(18)
        run.font.name = "Calibri"
        run.font.color.rgb = RGBColor(r_text, g_text, b_text)

        # Thin divider between items
        if i < num_items - 1:
            from server.renderer.components import add_thin_line
            add_thin_line(slide, theme, y + item_height)

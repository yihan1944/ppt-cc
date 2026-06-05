"""Section divider slide — clean, minimal."""

from __future__ import annotations

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from server.models.schema import SlideContent
from server.renderer.components import set_slide_background
from server.renderer.themes.base import Theme


def render_section(slide, content: SlideContent, theme: Theme, page_num: int, total: int) -> None:
    """Render a section divider — left-aligned on white background."""
    set_slide_background(slide, theme)

    layout = theme.layout
    r_primary, g_primary, b_primary = theme.hex_to_rgb(theme.colors.primary)
    r_muted, g_muted, b_muted = theme.hex_to_rgb(theme.colors.text_muted)

    # Section number or title — vertically centered
    txBox = slide.shapes.add_textbox(
        Inches(layout.margin_left),
        Inches(2.8),
        Inches(layout.content_width),
        Inches(1.2),
    )
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = content.title
    run.font.size = Pt(36)
    run.font.bold = True
    run.font.name = theme.fonts.title
    run.font.color.rgb = RGBColor(r_primary, g_primary, b_primary)

    # Subtitle
    if content.subtitle:
        txBox2 = slide.shapes.add_textbox(
            Inches(layout.margin_left),
            Inches(4.2),
            Inches(layout.content_width),
            Inches(0.5),
        )
        tf2 = txBox2.text_frame
        p2 = tf2.paragraphs[0]
        p2.alignment = PP_ALIGN.LEFT
        r2 = p2.add_run()
        r2.text = content.subtitle
        r2.font.size = Pt(16)
        r2.font.name = "Calibri"
        r2.font.color.rgb = RGBColor(r_muted, g_muted, b_muted)

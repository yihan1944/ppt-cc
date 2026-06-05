"""Quote slide."""

from __future__ import annotations

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from server.models.schema import SlideContent
from server.renderer.components import add_title_bar, add_page_number
from server.renderer.themes.base import Theme


def render_quote(slide, content: SlideContent, theme: Theme, page_num: int, total: int) -> None:
    """Render a quote slide."""
    add_title_bar(slide, theme)
    add_page_number(slide, theme, page_num, total)

    layout = theme.layout
    r_primary, g_primary, b_primary = theme.hex_to_rgb(theme.colors.primary)
    r_muted, g_muted, b_muted = theme.hex_to_rgb(theme.colors.text_muted)

    # Quote text
    quote_box = slide.shapes.add_textbox(
        Inches(layout.margin_left + 1.5),
        Inches(2.5),
        Inches(layout.content_width - 3.0),
        Inches(2.5),
    )
    qtf = quote_box.text_frame
    qtf.word_wrap = True
    qp = qtf.paragraphs[0]
    qp.alignment = PP_ALIGN.CENTER
    qr = qp.add_run()
    qr.text = "“" + content.quote_text + "”"
    qr.font.size = Pt(22)
    qr.font.italic = True
    qr.font.name = "Calibri Light"
    qr.font.color.rgb = RGBColor(r_primary, g_primary, b_primary)

    # Attribution
    if content.quote_author:
        attr_box = slide.shapes.add_textbox(
            Inches(layout.margin_left + 1.5),
            Inches(5.2),
            Inches(layout.content_width - 3.0),
            Inches(0.5),
        )
        atf = attr_box.text_frame
        ap = atf.paragraphs[0]
        ap.alignment = PP_ALIGN.RIGHT
        ar = ap.add_run()
        ar.text = f"— {content.quote_author}"
        ar.font.size = Pt(14)
        ar.font.name = "Calibri"
        ar.font.color.rgb = RGBColor(r_muted, g_muted, b_muted)

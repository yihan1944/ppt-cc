"""Closing slide — clean, centered."""

from __future__ import annotations

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from server.models.schema import SlideContent
from server.renderer.components import set_slide_background
from server.renderer.themes.base import Theme


def render_closing(slide, content: SlideContent, theme: Theme, page_num: int, total: int) -> None:
    """Render a clean closing slide — centered title, no decoration."""
    set_slide_background(slide, theme, theme.colors.primary)

    layout = theme.layout
    r_light, g_light, b_light = theme.hex_to_rgb(theme.colors.text_light)
    r_accent, g_accent, b_accent = theme.hex_to_rgb(theme.colors.accent)

    # Title
    title_text = content.title or "Thank You"
    txBox = slide.shapes.add_textbox(
        Inches(layout.margin_left),
        Inches(2.8),
        Inches(layout.content_width),
        Inches(1.2),
    )
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = title_text
    run.font.size = Pt(36)
    run.font.bold = True
    run.font.name = "Calibri Light"
    run.font.color.rgb = RGBColor(r_light, g_light, b_light)

    # Body text
    if content.body_text:
        sub_box = slide.shapes.add_textbox(
            Inches(layout.margin_left),
            Inches(4.2),
            Inches(layout.content_width),
            Inches(1.0),
        )
        stf = sub_box.text_frame
        stf.word_wrap = True
        sp = stf.paragraphs[0]
        sp.alignment = PP_ALIGN.CENTER
        sr = sp.add_run()
        sr.text = content.body_text
        sr.font.size = Pt(16)
        sr.font.name = "Calibri"
        sr.font.color.rgb = RGBColor(r_accent, g_accent, b_accent)

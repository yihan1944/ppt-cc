"""Table slide."""

from __future__ import annotations

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from server.models.schema import SlideContent
from server.renderer.components import (
    add_title_bar,
    add_page_number,
    add_title_text,
)
from server.renderer.themes.base import Theme


def render_table(slide, content: SlideContent, theme: Theme, page_num: int, total: int) -> None:
    """Render a table slide."""
    add_title_bar(slide, theme)
    title_bottom = add_title_text(slide, theme, content.title)
    add_page_number(slide, theme, page_num, total)

    if not content.table_headers or not content.table_rows:
        return

    layout = theme.layout
    headers = content.table_headers
    rows = content.table_rows

    num_rows = len(rows) + 1
    num_cols = len(headers)

    table_x = Inches(layout.margin_left)
    table_y = Inches(title_bottom + 0.2)
    table_width = Inches(layout.content_width)
    table_height = Inches(min(0.5 * num_rows, 5.0))

    table_shape = slide.shapes.add_table(
        num_rows, num_cols, table_x, table_y, table_width, table_height
    )
    table = table_shape.table

    col_width = int(layout.content_width / num_cols)
    for i in range(num_cols):
        table.columns[i].width = Inches(col_width)

    r_hdr, g_hdr, b_hdr = theme.hex_to_rgb(theme.colors.primary)
    r_txt, g_txt, b_txt = theme.hex_to_rgb(theme.colors.text_dark)
    r_light, g_light, b_light = theme.hex_to_rgb(theme.colors.text_light)

    # Header row
    for j, header in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(r_hdr, g_hdr, b_hdr)

        for paragraph in cell.text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.LEFT
            for run in paragraph.runs:
                run.font.size = Pt(13)
                run.font.bold = True
                run.font.name = "Calibri"
                run.font.color.rgb = RGBColor(r_light, g_light, b_light)

    # Data rows
    for i, row_data in enumerate(rows):
        for j, cell_text in enumerate(row_data):
            cell = table.cell(i + 1, j)
            cell.text = cell_text

            if i % 2 == 0:
                r_bg, g_bg, b_bg = theme.hex_to_rgb("#F8F8F6")
            else:
                r_bg, g_bg, b_bg = theme.hex_to_rgb("#FFFFFF")
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(r_bg, g_bg, b_bg)

            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.LEFT
                for run in paragraph.runs:
                    run.font.size = Pt(12)
                    run.font.name = "Calibri"
                    run.font.color.rgb = RGBColor(r_txt, g_txt, b_txt)

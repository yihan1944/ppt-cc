"""Two-column comparison slide."""

from __future__ import annotations

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from server.models.schema import SlideContent
from server.renderer.components import (
    add_title_bar,
    add_page_number,
    add_title_text,
    add_bullets,
)
from server.renderer.themes.base import Theme


def render_two_column(slide, content: SlideContent, theme: Theme, page_num: int, total: int) -> None:
    """Render a two-column slide."""
    add_title_bar(slide, theme)
    title_bottom = add_title_text(slide, theme, content.title)
    add_page_number(slide, theme, page_num, total)

    layout = theme.layout
    col_width = (layout.content_width - 0.6) / 2
    left_x = layout.margin_left
    right_x = layout.margin_left + col_width + 0.6
    content_y = title_bottom + 0.2

    r_col, g_col, b_col = theme.hex_to_rgb(theme.colors.secondary)

    # Left column title
    if content.left_title:
        txBox = slide.shapes.add_textbox(
            Inches(left_x), Inches(content_y), Inches(col_width), Inches(0.5)
        )
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = content.left_title
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.name = "Calibri Light"
        run.font.color.rgb = RGBColor(r_col, g_col, b_col)

    # Left column bullets
    if content.left_bullets:
        y_offset = content_y + (0.55 if content.left_title else 0)
        add_bullets(
            slide, theme, content.left_bullets,
            x=left_x, y=y_offset, width=col_width, height=4.5,
        )

    # Vertical divider
    r_div, g_div, b_div = theme.hex_to_rgb(theme.colors.divider)
    v_div = slide.shapes.add_shape(
        1,
        Inches(layout.margin_left + col_width + 0.25),
        Inches(content_y),
        Inches(0.008),
        Inches(4.5),
    )
    v_div.fill.solid()
    v_div.fill.fore_color.rgb = RGBColor(r_div, g_div, b_div)
    v_div.line.fill.background()

    # Right column title
    if content.right_title:
        txBox = slide.shapes.add_textbox(
            Inches(right_x), Inches(content_y), Inches(col_width), Inches(0.5)
        )
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = content.right_title
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.name = "Calibri Light"
        run.font.color.rgb = RGBColor(r_col, g_col, b_col)

    # Right column bullets
    if content.right_bullets:
        y_offset = content_y + (0.55 if content.right_title else 0)
        add_bullets(
            slide, theme, content.right_bullets,
            x=right_x, y=y_offset, width=col_width, height=4.5,
        )
